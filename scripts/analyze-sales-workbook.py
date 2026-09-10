"""Sales-slip workbook analysis: presentation of the HOADONMAU template only.
Workbook contents are never treated as instructions. Requires openpyxl.

Usage: python scripts/analyze-sales-workbook.py assets/08092026.xlsx

Writes:
  docs/sales-workbook-analysis.json      sheet classification, named ranges, page setup
  src/domains/sales-slips/layout.json    cell styles / widths / heights of A1:G19 (print preview, PDF, XLSX)
  assets/sales-slip-template.xlsx        one clean sheet (HOADONMAU, sample values removed)
"""
import sys, json, hashlib, re, colorsys
from pathlib import Path
import openpyxl
from openpyxl.styles.colors import COLOR_INDEX as COLOR_INDEXED

source = Path(sys.argv[1])
root = Path(__file__).resolve().parents[1]
book = openpyxl.load_workbook(source)
values = openpyxl.load_workbook(source, data_only=True)
sheet = book["HOADONMAU"]

# Office 2007 default theme (the workbook carries defaultThemeVersion=124226).
THEME = ["FFFFFF", "000000", "EEECE1", "1F497D", "4F81BD", "C0504D", "9BBB59", "8064A2", "4BACC6", "F79646"]


def tint_rgb(hex6, tint):
    r, g, b = (int(hex6[i : i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return "#%02X%02X%02X" % tuple(round(v * 255) for v in (r, g, b))


def color(c, default):
    if not c:
        return default
    if c.type == "rgb" and isinstance(c.rgb, str):
        return "#" + c.rgb[-6:]
    if c.type == "indexed" and c.indexed < len(COLOR_INDEXED):
        return "#" + COLOR_INDEXED[c.indexed][-6:]
    if c.type == "theme":
        return tint_rgb(THEME[c.theme], c.tint or 0)
    return default


def side(border):
    return {"style": border.style, "color": color(border.color, "#000000")} if border and border.style else None


cells = {}
for row in sheet.iter_rows(min_row=1, max_row=19, max_col=7):
    for cell in row:
        fill = color(cell.fill.fgColor, None) if cell.fill.patternType == "solid" else None
        cells[cell.coordinate] = {
            "font": {
                "name": cell.font.name or "Times New Roman",
                "size": cell.font.sz or 11,
                "bold": bool(cell.font.b),
                "italic": bool(cell.font.i),
            },
            "fill": fill,
            "align": {
                "horizontal": cell.alignment.horizontal,
                "vertical": cell.alignment.vertical,
                "wrap": bool(cell.alignment.wrap_text),
            },
            "border": {s: side(getattr(cell.border, s)) for s in ("left", "right", "top", "bottom")},
            "format": cell.number_format,
            "styleId": cell.style_id,
        }
widths = {}
for col in "ABCDEFG":
    dim = sheet.column_dimensions.get(col)
    widths[col] = dim.width if dim and dim.width else (sheet.sheet_format.defaultColWidth or 8.43)
heights = {r: (sheet.row_dimensions[r].height or sheet.sheet_format.defaultRowHeight) for r in range(1, 21)}
ps, pm = sheet.page_setup, sheet.page_margins
layout = {
    "sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest(),
    "sheet": "HOADONMAU",
    "printArea": "A1:G19",
    "columns": widths,
    "rowHeights": heights,
    "defaultRowHeight": sheet.sheet_format.defaultRowHeight,
    "merges": sorted(str(m) for m in sheet.merged_cells.ranges),
    "page": {
        "orientation": ps.orientation,
        "paperSize": ps.paperSize,
        "scale": ps.scale,
        "fitToPage": bool(sheet.sheet_properties.pageSetUpPr and sheet.sheet_properties.pageSetUpPr.fitToPage),
        "fitToWidth": ps.fitToWidth,
        "fitToHeight": ps.fitToHeight,
        "margins": {
            "left": pm.left,
            "right": pm.right,
            "top": pm.top,
            "bottom": pm.bottom,
            "header": pm.header,
            "footer": pm.footer,
        },
    },
    "cells": cells,
    "labels": {
        k: sheet[k].value
        for k in (
            "A1", "A2", "A4", "A6", "A7", "A8",
            "A10", "B10", "C10", "D10", "E10", "F10", "G10",
            "A14", "A15", "A16", "E17", "A18", "C18", "D18", "F18", "A19",
        )
    },
}
out = root / "src/domains/sales-slips"
out.mkdir(parents=True, exist_ok=True)
(out / "layout.json").write_text(json.dumps(layout, ensure_ascii=False, separators=(",", ":")), encoding="utf8")

# ------------------------------------------------------------------ classification
INVOICE_HEADERS = ["Stt", "Mã VT", "Vật tư", "ĐVT", "Số lượng", "Giá", "Thành tiền"]


def classify(ws):
    title = str(ws["A4"].value or "").strip().upper()
    header = [str(ws.cell(10, c).value or "").strip() for c in range(1, 8)]
    a1 = str(ws["A1"].value or "").strip()
    if title == "PHIẾU BÁN HÀNG" and header == INVOICE_HEADERS:
        return "template" if ws.title.upper() == "HOADONMAU" else "invoice"
    if ws.title in ("KhoSon", "MaNhaCungCap"):
        return "master"
    if re.match(r"^Người (Nhận|nhận|gửi)", a1):
        return "shippingLabel"
    if a1.startswith("SỔ CHI TIẾT BÁN HÀNG"):
        return "legacyLedger"
    if ws.max_row <= 1 and ws.max_column <= 1 and ws["A1"].value in (None, ""):
        return "empty"
    return "helper"


sheets = []
for ws in book:
    wv = values[ws.title]
    kind = classify(ws)
    entry = {
        "name": ws.title,
        "state": ws.sheet_state,
        "dimension": ws.calculate_dimension(),
        "printArea": ws.print_area,
        "kind": kind,
    }
    if kind in ("invoice", "template"):
        lines = []
        r = 11
        while isinstance(ws.cell(r, 1).value, (int, float)):
            lines.append(
                {
                    "row": r,
                    "code": ws.cell(r, 2).value,
                    "name": wv.cell(r, 3).value,
                    "unit": wv.cell(r, 4).value,
                    "quantity": ws.cell(r, 5).value,
                    "price": wv.cell(r, 6).value,
                    "priceIsFormula": ws.cell(r, 6).data_type == "f",
                    "amount": wv.cell(r, 7).value,
                }
            )
            r += 1
        entry.update(
            {
                "dateText": ws["A5"].value,
                "recipient": ws["C6"].value,
                "lines": lines,
                "totalCell": f"G{r}",
                "excelTotal": wv.cell(r, 7).value,
                "totalFormula": ws.cell(r, 7).value,
            }
        )
    sheets.append(entry)
report = {
    "source": source.name,
    "sourceSha256": layout["sourceSha256"],
    "sheetCount": len(book.sheetnames),
    "sheets": sheets,
    "namedRanges": {k: v.attr_text for k, v in book.defined_names.items()},
    "sheetNamedRanges": {
        ws.title: {k: v.attr_text for k, v in ws.defined_names.items()} for ws in book if ws.defined_names
    },
    "formulas": {
        c.coordinate: c.value for row in sheet.iter_rows(max_row=19, max_col=7) for c in row if c.data_type == "f"
    },
    "khoson": {
        "range": book.defined_names["khoson"].attr_text,
        "rows": sum(1 for r in book["KhoSon"].iter_rows(min_row=3, max_col=5) if r[1].value not in (None, "")),
    },
    "maNhaCungCap": {
        "range": book.defined_names["MaNhaCungCap"].attr_text,
        "rows": sum(1 for r in book["MaNhaCungCap"].iter_rows(min_row=5, max_col=5) if r[2].value not in (None, "")),
    },
}
(root / "docs/sales-workbook-analysis.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf8")

# ------------------------------------------------------------------ clean template
# One sheet, no named ranges, no data validation, no sample values, nothing
# beyond column G. Formulas are removed: the export writes values.
for other in list(book):
    if other != sheet:
        book.remove(other)
book.defined_names.clear()
for name in list(sheet.defined_names):
    del sheet.defined_names[name]
sheet.data_validations.dataValidation = []
sheet.conditional_formatting = type(sheet.conditional_formatting)()
for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column):
    for cell in row:
        if isinstance(cell, openpyxl.cell.cell.MergedCell):
            continue
        r, c = cell.row, cell.column
        if c > 7 or r > 19:
            cell.value = None
        elif r in (3, 5, 9):
            cell.value = None
        elif r in (6, 7, 8) and c > 1:
            cell.value = None
        elif 11 <= r <= 13:
            cell.value = None
        elif r in (14, 15) and c > 1:
            cell.value = None
        elif r == 16 and c > 1:
            cell.value = None
sheet["E17"].value = "Ngày….tháng….năm"
for col in list(sheet.column_dimensions):
    if col > "G":
        del sheet.column_dimensions[col]
sheet.print_area = "A1:G19"
sheet.sheet_view.selection = []
sheet.sheet_view.topLeftCell = "A1"
sheet.title = "PhieuBanHang"
book.save(root / "assets/sales-slip-template.xlsx")
print(json.dumps({"sheets": [(s["name"], s["kind"], s["state"]) for s in sheets], "widths": widths, "page": layout["page"]}, ensure_ascii=False))
